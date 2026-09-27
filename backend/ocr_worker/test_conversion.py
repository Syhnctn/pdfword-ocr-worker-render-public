"""Local end-to-end sanity checks for the PDF -> DOCX conversion pipeline.

Run: python test_conversion.py
"""
from __future__ import annotations

import base64
import io
import os
import pathlib
import sys
import tempfile
import time
import zipfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

import pymupdf
from docx import Document as DocxDocument

import main as worker


def build_one_page_sample() -> str:
    """A one-page text PDF written to a temp file; returns its path."""
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 100), "Tek sayfa icerik metni", fontsize=20)
    pdf_bytes = document.tobytes()
    document.close()

    handle = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    try:
        handle.write(pdf_bytes)
        handle.flush()
        return handle.name
    finally:
        handle.close()


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)
    print(f"PASS: {message}")


def build_sample_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)

    # H1-sized heading, body text, a bold sentence, a bullet list, and a
    # ruled table that find_tables() can detect.
    page.insert_text((72, 90), "Quarterly Report", fontname="hebo", fontsize=24)
    page.insert_text(
        (72, 130),
        "Overall the team delivered strong results this quarter.",
        fontname="helv",
        fontsize=11,
    )
    page.insert_text(
        (72, 155),
        "Revenue grew by 18% across all regions.",
        fontname="hebo",
        fontsize=11,
    )
    y = 190
    for item in ("First milestone completed", "Second milestone in progress"):
        page.insert_text((72, y), f"• {item}", fontname="helv", fontsize=11)
        y += 20

    top = y + 20
    row_height = 22
    xs = [72, 220, 360, 500]
    rows = (
        ("Region", "Growth", "Notes"),
        ("EMEA", "12%", "stable"),
        ("APAC", "24%", "strong"),
    )
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            page.insert_text(
                (xs[column_index] + 4, top + row_index * row_height + 15),
                value,
                fontname="helv",
                fontsize=10,
            )
    for row_index in range(len(rows) + 1):
        page.draw_line(
            (xs[0], top + row_index * row_height),
            (xs[3], top + row_index * row_height),
        )
    for x in xs:
        page.draw_line((x, top), (x, top + len(rows) * row_height))

    # Second page so page breaks are exercised.
    page2 = document.new_page(width=595, height=842)
    page2.insert_text((72, 90), "Appendix", fontname="hebo", fontsize=18)
    page2.insert_text(
        (72, 200),
        "Supporting details for the report.",
        fontname="helv",
        fontsize=11,
    )

    data = document.tobytes()
    document.close()
    return data


def docx_xml(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml")


def main() -> int:
    pdf_bytes = build_sample_pdf()

    sections = worker.extract_pdf_text_sections(pdf_bytes)
    check(len(sections) == 2, "extracts both pages structurally")
    page1, page2 = sections[0][1], sections[1][1]

    check("# Quarterly Report" in page1, "large bold text becomes an H1 heading")
    check(
        "**Revenue grew by 18% across all regions.**" in page1,
        "bold sentence keeps emphasis markers",
    )
    check(
        "- First milestone completed" in page1
        and "- Second milestone in progress" in page1,
        "bullet lines become markdown list items",
    )
    check(
        "| Region | Growth | Notes |" in page1 and "| APAC | 24% | strong |" in page1,
        "ruled table is extracted as a markdown table",
    )
    check("# Appendix" in page2, "second page heading is detected")

    markdown = worker.build_markdown_from_extracted_files(
        [{"name": "ornek.pdf", "sections": sections}]
    )
    check(markdown.startswith("# ornek.pdf"), "file name becomes the document title")
    check(
        worker._PAGEBREAK_MARKER in markdown,
        "pages are separated with a page break marker",
    )
    check(
        worker.build_markdown_from_extracted_files([]) == "",
        "empty input yields empty markdown (placeholder path)",
    )
    noted = worker.build_markdown_from_extracted_files(
        [{"name": "a.pdf", "sections": [(1, "Hello")], "note": "Extracted with OCR."}]
    )
    check("> Extracted with OCR." in noted, "worker notes render as a blockquote")

    docx_bytes = worker.markdown_to_docx_bytes(markdown)
    document = DocxDocument(io.BytesIO(docx_bytes))
    styles = [paragraph.style.name for paragraph in document.paragraphs]
    check(
        any(style.startswith("Heading") for style in styles),
        "docx contains real heading styles",
    )
    check("List Bullet" in styles, "docx contains bullet list style")
    check(len(document.tables) == 1, "docx contains one real Word table")
    cell_values = [
        cell.text for row in document.tables[0].rows for cell in row.cells
    ]
    check(
        "Region" in cell_values and "APAC" in cell_values,
        "table cells keep their values",
    )
    bold_runs = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.bold
    ]
    check(
        any("Revenue grew" in run.text for run in bold_runs),
        "bold text becomes bold runs",
    )
    check(
        b'type="page"' in docx_xml(docx_bytes),
        "page break marker becomes a real Word page break",
    )

    legacy = worker.markdown_to_docx_bytes("Intro\n\n### Page 2\n\nBody")
    check(
        b'type="page"' in docx_xml(legacy),
        "legacy '### Page N' headings still map to page breaks",
    )
    # Direct bytes conversion endpoint (used by the app as a fallback when
    # signed storage uploads are unavailable). Runs without Supabase env, so
    # the DOCX must come back inline as base64.
    from fastapi.testclient import TestClient

    client = TestClient(worker.app)
    response = client.post(
        "/internal/convert",
        files=[
            (
                "files",
                ("sample_report.pdf", pdf_bytes, "application/pdf"),
            )
        ],
        data={"source": "test"},
    )
    check(response.status_code == 200, "direct convert endpoint responds 200")
    payload = response.json()
    check(payload.get("status") == "succeeded", "direct convert reports success")
    check(
        bool(payload.get("docx_base64")),
        "direct convert returns inline docx when supabase is not configured",
    )
    inline_docx = base64.b64decode(payload["docx_base64"])
    check(
        b"Quarterly Report" in docx_xml(inline_docx),
        "inline docx contains the extracted heading",
    )

    original_fitz = worker.fitz


    original_fitz = worker.fitz
    try:
        worker.fitz = None
        fallback = worker.extract_pdf_text_sections(pdf_bytes)
        check(bool(fallback), "falls back to pypdf when PyMuPDF is unavailable")
        check(
            "Quarterly Report" in fallback[0][1],
            "pypdf fallback returns the raw text",
        )
    finally:
        worker.fitz = original_fitz

    # --- chunked OCR (page-group subprocesses) -------------------------
    check(worker.ocr_chunk_pages() == 1, "default chunk size is 1 page")

    multi_page = pymupdf.open()
    for index in range(5):
        page = multi_page.new_page(width=595, height=842)
        page.insert_text((72, 100), f"Sayfa {index + 1} icerik metni", fontsize=20)
    multi_pdf = multi_page.tobytes()
    multi_page.close()

    with tempfile.TemporaryDirectory() as work_dir:
        source = pathlib.Path(work_dir) / "source.pdf"
        source.write_bytes(multi_pdf)
        sizes: list[int] = []
        for page_index in range(5):
            chunk = pathlib.Path(work_dir) / f"page-{page_index}.pdf"
            written = worker._write_pdf_slice(str(source), str(chunk), page_index, 1)
            sizes.append(pymupdf.open(chunk).page_count if written else -1)
        check(sizes == [1, 1, 1, 1, 1], "5 pages split into five single-page files")

    chunk_result = worker._run_ocr_subprocess(build_one_page_sample(), 120.0)
    check(bool(chunk_result.get("ok")), "isolated OCR subprocess returns a result")
    check(
        len(chunk_result.get("sections") or []) == 1,
        "isolated subprocess reports its single page",
    )
    check(
        [int(item[0]) for item in (chunk_result.get("sections") or [])] == [1],
        "isolated subprocess keeps 1-based page numbers",
    )

    merged_sections, merged_note = worker.extract_sections_via_chunked_subprocess(
        multi_pdf
    )
    check(
        [int(index) for index, _ in merged_sections] == [1, 2, 3, 4, 5],
        "chunked path reassembles every page in order",
    )
    check(
        all("Sayfa" in text for _, text in merged_sections),
        "chunked path keeps the text of each page",
    )

    progress_calls: list[tuple[int, int]] = []
    progressed_sections, _ = worker.extract_sections_via_chunked_subprocess(
        multi_pdf,
        on_page_done=lambda group, total: progress_calls.append((group, total)),
    )
    check(
        len(progress_calls) == 5,
        "chunked path reports progress once per page",
    )
    check(
        [int(index) for index, _ in progressed_sections] == [1, 2, 3, 4, 5],
        "progress reporting does not disturb the reassembled pages",
    )

    # --- background job runner -----------------------------------------
    check(worker.background_process_spawned(), "background jobs are enabled")

    original_background = os.environ.get("OCR_BACKGROUND_JOBS")
    try:
        os.environ["OCR_BACKGROUND_JOBS"] = "false"
        check(
            not worker.background_process_spawned(),
            "OCR_BACKGROUND_JOBS=false switches back to inline processing",
        )
    finally:
        if original_background is None:
            os.environ.pop("OCR_BACKGROUND_JOBS", None)
        else:
            os.environ["OCR_BACKGROUND_JOBS"] = original_background

    os.environ["OCR_BACKGROUND_JOBS"] = "true"
    started = time.monotonic()
    dispatched = worker.process_job_safely("background-test-job")
    elapsed = time.monotonic() - started
    check(
        dispatched.get("status") == "processing",
        "background dispatch reports the job as processing",
    )
    check(
        dispatched.get("mode") == "background",
        "background dispatch marks the response as background mode",
    )
    check(
        elapsed < 5.0,
        "background dispatch returns immediately without waiting for OCR",
    )

    print("ALL CONVERSION CHECKS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
